#!/usr/bin/env python3
"""Build editable BMC manuscript and supplement DOCX files from audited sources."""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "BMC_Bioinformatics_submission"
MASTER = ROOT / "results/bmc_master"
ASSETS = OUT / "assets"

REF_LABELS = {
    "fig:framework": "1",
    "fig:audit": "2",
    "fig:matched": "3",
    "fig:probes": "4",
    "tab:datasets": "1",
    "tab:splits": "2",
    "tab:matched": "3",
    "tab:probes": "4",
    "tab:checklist": "5",
}


def shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def set_width(cell, width_inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None, font_size: float = 8.2) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    keep_row_together(table.rows[0])
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = text
        shade(cell, "D9EAF7")
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if widths: set_width(cell, widths[i])
    for row in rows:
        tr = table.add_row()
        keep_row_together(tr)
        cells = tr.cells
        for i, text in enumerate(row):
            cells[i].text = str(text)
            if widths: set_width(cells[i], widths[i])
            for p in cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(font_size)
    doc.add_paragraph()


def basic_doc(title: str, authors: str | None = None) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.75); section.bottom_margin = Inches(.75)
    section.left_margin = Inches(.82); section.right_margin = Inches(.82)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"; normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size, color in [("Title", 16, "17365D"), ("Heading 1", 13, "17365D"), ("Heading 2", 11, "2F5597")]:
        st = styles[style_name]
        st.font.name = "Arial"; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(10); st.paragraph_format.space_after = Pt(4)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)
    if authors:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(authors); r.font.size = Pt(10); r.font.italic = True
    return doc


def read_bib(path: Path) -> dict[str, dict[str, str]]:
    text = path.read_text(encoding="utf-8")
    entries = {}
    for key, body in re.findall(r"@\w+\{([^,]+),(.*?)(?=\n@\w+\{|\Z)", text, re.S):
        fields = {}
        for name, value in re.findall(r"(\w+)\s*=\s*\{((?:[^{}]|\{[^{}]*\})*)\}", body, re.S):
            fields[name.lower()] = re.sub(r"\s+", " ", value.strip())
        entries[key] = fields
    return entries


def citation_mapping(tex: str) -> dict[str, int]:
    values = []
    for match in re.finditer(r"\\cite[pt]?\{([^}]+)\}", tex):
        values.extend(x.strip() for x in match.group(1).split(","))
    mapping = {}
    for value in values:
        if value not in mapping:
            mapping[value] = len(mapping) + 1
    return mapping


def clean(text: str, citations: dict[str, int]) -> str:
    def replace_cite(match):
        nums = sorted(citations[x.strip()] for x in match.group(1).split(",") if x.strip() in citations)
        return "[" + ", ".join(str(x) for x in nums) + "]"
    def replace_ref(match):
        return REF_LABELS.get(match.group(1), "")
    text = re.sub(r"\\cite[pt]?\{([^}]+)\}", replace_cite, text)
    text = re.sub(r"\\ref\{([^}]+)\}", replace_ref, text)
    text = re.sub(r"\\url\{([^}]*)\}", r"\1", text)
    text = re.sub(r"\\texttt\{([^}]*)\}", lambda m: "`" + m.group(1).replace(r"\_", "_") + "`", text)
    text = re.sub(r"\\(?:emph|textbf)\{([^}]*)\}", r"\1", text)
    text = text.replace(r"\%", "%").replace(r"\_", "_").replace("~", " ").replace("---", "—").replace("--", "–")
    text = text.replace(r"\pm", "±").replace(r"\alpha", "α").replace(r"\times", "×").replace(r"\textwidth", "text width")
    text = text.replace("1 × 10^{-3}", "1 × 10⁻³").replace("10^{-3}", "10⁻³")
    text = text.replace("$", "")
    text = text.replace("r_i=(c_i,d_i,s_i,z_i,m_i,y_i)", "rᵢ = (cᵢ, dᵢ, sᵢ, zᵢ, mᵢ, yᵢ)")
    for raw, pretty in {
        "r_i": "rᵢ",
        "c_i": "cᵢ",
        "d_i": "dᵢ",
        "s_i": "sᵢ",
        "z_i": "zᵢ",
        "m_i": "mᵢ",
        "y_i": "yᵢ",
    }.items():
        text = text.replace(raw, pretty)
    text = re.sub(r"\\[a-zA-Z]+(?:\[[^]]*\])?", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def add_clean_paragraphs(doc: Document, block: str, citations: dict[str, int]) -> None:
    for para in re.split(r"\n\s*\n", block.strip()):
        text = clean(para, citations)
        if text:
            doc.add_paragraph(text)


def bib_plain(text: str) -> str:
    accents = {
        r"\'a": "á", r"\'e": "é", r"\'i": "í", r"\'o": "ó", r"\'u": "ú",
        r"\'A": "Á", r"\'E": "É", r"\'I": "Í", r"\'O": "Ó", r"\'U": "Ú",
        r'\"a': "ä", r'\"e': "ë", r'\"i': "ï", r'\"o': "ö", r'\"u': "ü",
        r'\"A': "Ä", r'\"E': "Ë", r'\"I': "Ï", r'\"O': "Ö", r'\"U': "Ü",
        r"\~n": "ñ", r"\~N": "Ñ",
    }
    for raw, cooked in accents.items():
        text = text.replace(raw, cooked)
    text = re.sub(r"\{\\['\"]([A-Za-z])\}", lambda m: accents.get(r"\'" + m.group(1), m.group(1)), text)
    text = text.replace("{", "").replace("}", "").replace("--", "–")
    return text


def caption_from(block: str, citations: dict[str, int]) -> str:
    match = re.search(r"\\caption\{(.*?)\}", block, re.S)
    return clean(match.group(1), citations) if match else ""


def insert_figure(doc: Document, filename: str, caption: str, prefix: str | None = None) -> None:
    path = ASSETS / "figures" / filename
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(6.45))
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_prefix = prefix if prefix is not None else "Figure " + filename.split("_")[1]
    r = p.add_run(figure_prefix + ". " + caption)
    r.italic = True; r.font.size = Pt(8.6)


def extracted_blocks(tex: str):
    body = tex[tex.index(r"\section{Background}"):]
    body = body.split(r"\bibliographystyle", 1)[0]
    pattern = r"(\\section\*?\{.*?\}|\\subsection\{.*?\}|\\begin\{figure\*?\}.*?\\end\{figure\*?\}|\\begin\{table\*?\}.*?\\end\{table\*?\})"
    for part in re.split(pattern, body, flags=re.S):
        if part.strip():
            yield part.strip()


def table_from_label(doc: Document, block: str) -> None:
    if "tab:datasets" in block:
        d = pd.read_csv(MASTER / "master_dataset_characteristics.csv")
        fmt_int = lambda v: f"{int(v):,}"
        rows = [[x.dataset, x.stage, fmt_int(x.records), fmt_int(x.unique_drugs), fmt_int(x.unique_cell_contexts), "NA" if pd.isna(x.unique_scaffolds_nonmissing) else fmt_int(x.unique_scaffolds_nonmissing), fmt_int(x.unique_doses)] for _, x in d.iterrows()]
        add_table(doc, ["Dataset", "Stage", "Records", "Drugs", "Cells", "Scaffolds", "Doses"], rows, [1.05, 1.55, .55, .5, .45, .62, .45])
    elif "tab:splits" in block:
        s = pd.read_csv(MASTER / "master_split_summary.csv")
        s = s[s.count_stage.eq("canonical_manuscript_reporting")]
        rows = []
        for (dataset, split), g in s.groupby(["dataset", "split"]):
            x = g.set_index("quantity")
            val = lambda key: f"{x.loc[key, 'mean']:,.1f} ± {x.loc[key, 'sd']:.1f}"
            rows.append([dataset, x.iloc[0].split_label, int(x.loc["manifest_test", "n_seeds"]), val("manifest_train"), val("manifest_test"), val("manifest_excluded"), val("evaluable_test")])
        add_table(doc, ["Dataset", "Split", "Seeds", "Manifest train", "Manifest test", "Excluded", "Evaluable test"], rows, [.9, .95, .42, .92, .9, .75, .9], 7.6)
    elif "tab:matched" in block:
        c = pd.read_csv(MASTER / "master_contrast_summary.csv")
        c = c[c.comparison.eq("train_test_size_matched_random minus joint held-out")]
        names = {"ridge_cell_drug_fp": "Cell-context + fingerprint ridge", "control_conditioned_nonlinear_benchmark": "Control-conditioned nonlinear benchmark"}
        rows = [[x.dataset, names[x.model], int(x.n_paired_seeds), f"{x.mean_difference:.3f}", f"{x.bootstrap_95ci_low:.3f} to {x.bootstrap_95ci_high:.3f}"] for _, x in c.iterrows()]
        add_table(doc, ["Dataset", "Probe", "Seeds", "Difference", "95% CI"], rows, [1.0, 2.2, .55, .72, 1.35])
    elif "tab:probes" in block:
        p = pd.read_csv(MASTER / "master_performance_summary.csv")
        models = ["global_train_mean", "cell_context_mean", "drug_mean", "ridge_cell", "ridge_drug_fp", "ridge_cell_drug_fp", "ridge_cell_dose_drug_fp", "control_conditioned_nonlinear_benchmark"]
        names = {"global_train_mean": "Global mean", "cell_context_mean": "Cell-context or cell-line mean", "drug_mean": "Drug mean", "ridge_cell": "Cell-context ridge", "ridge_drug_fp": "Drug-fingerprint ridge", "ridge_cell_drug_fp": "Cell-context + drug-fingerprint ridge", "ridge_cell_dose_drug_fp": "Cell-line + dose + drug-fingerprint ridge", "control_conditioned_nonlinear_benchmark": "Control-conditioned nonlinear benchmark"}
        rows = []
        for dataset in ["OpenProblems", "Sci-Plex 3"]:
            for model in models:
                q = p[(p.dataset.eq(dataset)) & p.model.eq(model) & p.metric.eq("rowwise_pearson")].set_index("split_label")
                if not {"random", "scaffold held-out", "joint held-out"}.issubset(q.index): continue
                rows.append([dataset, names[model], *[f"{q.loc[v, 'mean']:.3f} ± {q.loc[v, 'sd']:.3f}" for v in ["random", "scaffold held-out", "joint held-out"]]])
        add_table(doc, ["Dataset", "Probe", "Random", "Scaffold held-out", "Joint held-out"], rows, [.82, 2.25, .9, .9, .9], 7.3)
    elif "tab:checklist" in block:
        rows = [
            ["Evaluation target", "State whether the task is local interpolation, chemical structural extrapolation, cellular extrapolation, or a joint task."],
            ["Manifest accounting", "Release fixed seed-specific train, test, and excluded identifiers and report their counts."],
            ["Neighborhood audit", "Report overlapping drug, scaffold, cell, dose, MoA, missingness, and nearest training-drug similarity fields where applicable."],
            ["Matched controls", "State exactly which training, test, and composition variables are matched."],
            ["Model contract", "Report inputs, preprocessing, train-only validation, initialization, and hyperparameter-selection restrictions."],
            ["Metric context", "Pair absolute response-vector fit with baseline-adjusted, residualized, DE-gene, pathway, or retrieval metrics when perturbation-specific signal is the target."],
            ["Result portability", "Release seed-level metrics, paired contrasts, alignment information, source data, and integrity checks."],
        ]
        add_table(doc, ["Item", "Minimum report"], rows, [1.35, 5.2], 8.5)


def add_references(doc: Document, mapping: dict[str, int], bib: dict[str, dict[str, str]]) -> None:
    def initials(given: str) -> str:
        vals = []
        for part in re.split(r"[\s\-]+", bib_plain(given).strip()):
            if part:
                vals.append(part[0].upper() + ".")
        return " ".join(vals)

    def format_author(raw: str) -> str:
        raw = bib_plain(raw.strip())
        if raw.lower() == "others":
            return "et al."
        if "," in raw:
            family, given = [x.strip() for x in raw.split(",", 1)]
            return (family + " " + initials(given)).strip()
        parts = raw.split()
        if len(parts) <= 1:
            return raw
        return (parts[-1] + " " + initials(" ".join(parts[:-1]))).strip()

    def format_authors(raw: str) -> str:
        raw_authors = [a.strip() for a in raw.split(" and ") if a.strip()]
        authors = [a for a in raw_authors if a.lower() != "others"]
        shown = authors[:3]
        text = ", ".join(format_author(a) for a in shown)
        if len(authors) > 3 or any(a.lower() == "others" for a in raw_authors):
            text += ", et al."
        return text

    def end_sentence(value: str) -> str:
        value = bib_plain(value).strip()
        if not value:
            return ""
        return value if value.endswith((".", "?", "!")) else value + "."

    doc.add_heading("References", level=1)
    for key, number in sorted(mapping.items(), key=lambda x: x[1]):
        x = bib.get(key, {})
        author = format_authors(x.get("author", ""))
        title = bib_plain(x.get("title", ""))
        source = bib_plain(x.get("journal", x.get("booktitle", x.get("publisher", ""))))
        year = x.get("year", "")
        volume = x.get("volume", "")
        issue = x.get("number", "")
        pages = bib_plain(x.get("pages", ""))
        doi = x.get("doi", "")
        source_bits = []
        if source:
            source_bits.append(source)
        if year:
            source_bits.append(year)
        if volume:
            source_bits.append("vol. " + volume + (f"({issue})" if issue else ""))
        if pages:
            source_bits.append("pp. " + pages)
        pieces = [end_sentence(author), end_sentence(title), end_sentence(", ".join(source_bits))]
        if doi:
            pieces.append("doi: " + doi + ".")
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Inches(.25); p.paragraph_format.first_line_indent = Inches(-.25)
        p.add_run(f"[{number}] " + " ".join(y for y in pieces if y))


def build_main() -> None:
    tex = (OUT / "manuscript_BMC_Bioinformatics.tex").read_text(encoding="utf-8")
    citations = citation_mapping(tex); bib = read_bib(OUT / "references.bib")
    title = clean(re.search(r"\\title\{(.*?)\}", tex, re.S).group(1), citations)
    doc = basic_doc(title, "Da Lin¹, Ying Chen², Yue Liu², Yu Zhang¹*")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("¹The Second Affiliated Hospital of Wenzhou Medical University, Wenzhou, Zhejiang, China\n²Wenzhou Medical University, Wenzhou, Zhejiang, China\n*Correspondence: zhangyu1@wzhealth.com").font.size = Pt(8.6)
    abstract = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", tex, re.S).group(1)
    doc.add_heading("Abstract", level=1)
    for part in re.split(r"\\textbf\{([^}]+)\}", abstract):
        text = clean(part, citations)
        if not text: continue
        if text.endswith(":"):
            p = doc.add_paragraph(); p.add_run(text).bold = True
        else:
            doc.add_paragraph(text)
    keywords = re.search(r"\\noindent\\textbf\{Keywords:\}\s*(.*)", tex).group(1)
    p = doc.add_paragraph(); p.add_run("Keywords: ").bold = True; p.add_run(clean(keywords, citations))
    table_counter = 0
    for block in extracted_blocks(tex):
        if block.startswith(r"\section"):
            title = clean(re.search(r"\{(.*?)\}", block, re.S).group(1), citations)
            doc.add_heading(title, level=1)
        elif block.startswith(r"\subsection"):
            title = clean(re.search(r"\{(.*?)\}", block, re.S).group(1), citations)
            doc.add_heading(title, level=2)
        elif block.startswith(r"\begin{figure"):
            match = re.search(r"\\includegraphics\[.*?\]\{(.*?)\}", block)
            if match: insert_figure(doc, match.group(1), caption_from(block, citations))
        elif block.startswith(r"\begin{table"):
            table_counter += 1
            if "tab:checklist" in block:
                doc.add_page_break()
            caption = caption_from(block, citations)
            p = doc.add_paragraph(); p.add_run(f"Table {table_counter}. " + caption).bold = True
            table_from_label(doc, block)
        else:
            add_clean_paragraphs(doc, block, citations)
    add_references(doc, citations, bib)
    doc.save(OUT / "manuscript_BMC_Bioinformatics.docx")


def build_supplement() -> None:
    doc = basic_doc("Supplementary material", "A leakage-aware benchmark for evaluating chemical and cellular generalization in single-cell perturbation-response prediction")
    doc.add_heading("Supplementary methods", level=1)
    doc.add_heading("Fixed manifests and integrity checks", level=2)
    doc.add_paragraph("For each seed, fixed manifests assign every record to train, test, or, for joint tasks, excluded. Random-control draws are deterministic functions of the same seed and fixed random manifest. Joint record identities are used only to obtain target counts and are never transferred into random-control partitions. Prediction archives are checked for aligned record identifiers, target and prediction vector length, and exclusion of test or excluded records from fitting, normalization, validation, early stopping, or hyperparameter selection.")
    doc.add_heading("Model input contracts", level=2)
    doc.add_paragraph("Transparent probes use only information available under the manifest. The control-conditioned nonlinear benchmark uses basal/control expression, a 1,024-bit fingerprint, log dose, and cell-line one-hot features; its validation partition is derived only from training records. It is not a full reproduction of TranSiGen. PRnet-interface results remain adaptation-status outputs.")
    doc.add_heading("Supplementary Figure S1", level=1)
    insert_figure(doc, "Supplementary_Figure_S1_nonlinear_training_loss.png", "Training and validation loss traces for the control-conditioned nonlinear benchmark under Sci-Plex 3 matched-random controls. Lines show the mean across completed fixed seeds and shaded regions show one standard deviation.", "Supplementary Figure S1")
    doc.add_heading("Supplementary Table S1. Fixed split accounting", level=1)
    table_from_label(doc, "tab:splits")
    doc.add_heading("Supplementary Table S2. Completed multi-probe performance", level=1)
    table_from_label(doc, "tab:probes")
    doc.add_heading("Supplementary Table S3. Train-and-test-size-matched random contrasts", level=1)
    table_from_label(doc, "tab:matched")
    doc.save(OUT / "supplementary_BMC_Bioinformatics.docx")


if __name__ == "__main__":
    build_main()
    build_supplement()
